import json, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
data=json.loads((ROOT/'content.json').read_text())
d=Document()
for border in list(d.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec=d.sections[0]
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=d.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(6)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(27)
d.styles['Heading 1'].font.size=Pt(23)
d.styles['Heading 2'].font.size=Pt(12)
d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
d.styles['Heading 2'].paragraph_format.space_after=Pt(3)
d.styles['Subtitle'].font.size=Pt(11)
h=sec.header.paragraphs[0]; h.text='OZEMPICAI   PRODUCT STRATEGY'
h.runs[0].font.size=Pt(8); h.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]; footer.text='Leadership proposal   •   September 2026'
footer.runs[0].font.size=Pt(8)
footer.add_run(' '*7+'Page ').font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for i,page in enumerate(data['pages']):
    if i: d.add_page_break()
    d.add_paragraph(page['title'], 'Title' if i==0 else 'Heading 1')
    if page.get('subtitle'): d.add_paragraph(page['subtitle'],'Subtitle')
    for heading,body in page['sections']:
        d.add_paragraph(heading,'Heading 2')
        p=d.add_paragraph(body)
        if i==11:
            p.paragraph_format.line_spacing=1.0
            for r in p.runs:r.font.size=Pt(9)
d.core_properties.title='OzempicAI product strategy and implementation'
d.core_properties.subject='Feature proposals, architecture, experiments and rollout'
d.core_properties.author='Product Strategy'
d.save(ROOT.parent/'deliverables/OzempicAI_Product_Strategy.docx')
print('Saved DOCX')
